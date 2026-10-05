"""Authorized write-only XLSX representation of all S13 CSV records."""
from __future__ import annotations
import csv, hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
HERE=Path(__file__).resolve().parent; D=HERE/'s13_dataset'; OUT=D/'S13_empirical_full_packet_streamed.xlsx'
FILES=[('Run index','s13_runs.csv'),('Events','s13_events.csv'),('Decisions','s13_decisions.csv'),('Receiver log','s13_receiver_log.csv'),('Receiver transitions','s13_receiver_transitions.csv'),('Outcomes','s13_outcomes.csv'),('Packets','s13_packets.csv')]
def main():
 wb=Workbook(write_only=True); head=PatternFill('solid',fgColor='1F4E78')
 readme=wb.create_sheet('Read me'); readme.freeze_panes='A2'; readme.column_dimensions['A'].width=25; readme.column_dimensions['B'].width=110; readme.row_dimensions[1].height=32; first=[]
 for value in ['S13 full packet export','Current evidence-bound empirical software-testbed package']:
  c=WriteOnlyCell(readme,value); c.font=Font(bold=True,color='FFFFFF'); c.fill=head; c.alignment=Alignment(wrap_text=True); first.append(c)
 readme.append(first); readme.append(['Scope','Packet/protocol observations only; no physical clock recovery, GNSS, SyncE, timing quality, or deployment claim.']); readme.append(['Precision','Capture timestamps, identities, hashes, frame bytes and original log text are stored as strings; do not coerce to numeric values.']); readme.append(['Linkage','Use run_id to join sheets. Original audited TIMESAFE workbook remains separate at outputs/manual_dataset_combined/ORAN_All_Current_Datasets_FINAL_PCAP.xlsx.']); readme.append(['Post-seal decisions','Decisions identifies SEALED evaluation evidence separately from retained post-seal rows.'])
 for title,file in FILES:
  ws=wb.create_sheet(title); ws.freeze_panes='A2'
  with (D/file).open(encoding='utf-8',newline='') as f:
   r=csv.reader(f); hdr=next(r); ws.row_dimensions[1].height=32
   for i,name in enumerate(hdr,1): ws.column_dimensions[get_column_letter(i)].width=42 if name in ('frame_hex','original_line') else 24 if 'identity' in name or 'sha256' in name else 16
   cells=[]
   for value in hdr:
    cell=WriteOnlyCell(ws,value); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=head; cell.alignment=Alignment(wrap_text=True); cells.append(cell)
   ws.append(cells)
   rows=1
   for row in r: ws.append(row); rows+=1
  ws.auto_filter.ref=f'A1:{get_column_letter(len(hdr))}{rows}'
 wb.save(OUT)
 print(hashlib.sha256(OUT.read_bytes()).hexdigest())
if __name__=='__main__': main()
