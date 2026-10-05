"""Authorized write-only XLSX representation of all S14 CSV records."""
from __future__ import annotations
import csv, hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
D = HERE / 's14_dataset'
OUT = D / 'S14_empirical_full_packet_streamed.xlsx'
FILES = [
    ('Run index', 's14_runs.csv'),
    ('Events', 's14_events.csv'),
    ('Decisions', 's14_decisions.csv'),
    ('Receiver log', 's14_receiver_log.csv'),
    ('Receiver transitions', 's14_receiver_transitions.csv'),
    ('Outcomes', 's14_outcomes.csv'),
    ('Packets', 's14_packets.csv')
]


def main():
    wb = Workbook(write_only=True)
    head = PatternFill('solid', fgColor='1F4E78')
    readme = wb.create_sheet('Read me')
    readme.freeze_panes = 'A2'
    readme.column_dimensions['A'].width = 25
    readme.column_dimensions['B'].width = 110
    readme.row_dimensions[1].height = 32
    first = []
    for value in ['S14 full packet export', 'Current evidence-bound empirical software-testbed package']:
        c = WriteOnlyCell(readme, value)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
        c.alignment = Alignment(wrap_text=True)
        first.append(c)
    readme.append(first)
    readme.append(['Scope', 'Packet/protocol observations only; no physical clock recovery, GNSS, SyncE, timing quality, or deployment claim.'])
    readme.append(['Precision', 'Capture timestamps, identities, hashes, frame bytes and original log text are stored as strings; do not coerce to numeric values.'])
    readme.append(['Linkage', 'Use run_id to join sheets across runs, events, decisions, receiver transitions, and packets.'])
    readme.append(['Post-seal decisions', 'Decisions identifies SEALED evaluation evidence separately from retained post-seal rows. S14 runs have zero post-seal lines.'])
    readme.append(['S14 evaluation', 'Evaluated against S14_BROADER_EXPERIMENT_PROTOCOL.json across 5 conditions with matched controls. These are DEVELOPMENT-SET results: the receiver-side dispersion statistic and its 29000 ns boundary were derived from S14, so S14 cannot validate them.'])
    readme.append(['Independent validation', 'NOT COMPLETE. The protocol is frozen at S15_INDEPENDENT_VALIDATION_PROTOCOL_V3.json and ZERO S15 runs have been executed. No S15 agreement, sensitivity or specificity figure exists.'])
    readme.append(['What is supported', 'Agreement between a packet-level decision and a receiver-side path-delay dispersion measurement in a software testbed, once S15 has run. NOT harmful-versus-harmless classification. Harmfulness remains UNKNOWN: no service requirement has been specified against which a path-delay dispersion could be judged harmful.'])
    readme.append(['Authoritative status', 'The Current status table in REMAINING_WORK_ACCEPTANCE_REGISTER.md governs. Earlier status statements there are marked superseded history.'])

    for title, file in FILES:
        ws = wb.create_sheet(title)
        ws.freeze_panes = 'A2'
        with (D / file).open(encoding='utf-8', newline='') as f:
            r = csv.reader(f)
            hdr = next(r)
            ws.row_dimensions[1].height = 32
            for i, name in enumerate(hdr, 1):
                ws.column_dimensions[get_column_letter(i)].width = 42 if name in ('frame_hex', 'original_line') else 24 if 'identity' in name or 'sha256' in name else 16
            cells = []
            for value in hdr:
                cell = WriteOnlyCell(ws, value)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = head
                cell.alignment = Alignment(wrap_text=True)
                cells.append(cell)
            ws.append(cells)
            rows = 1
            for row in r:
                ws.append(row)
                rows += 1
            ws.auto_filter.ref = f'A1:{get_column_letter(len(hdr))}{rows}'
    wb.save(OUT)
    print("XLSX_SAVED:", hashlib.sha256(OUT.read_bytes()).hexdigest())


if __name__ == '__main__':
    main()
