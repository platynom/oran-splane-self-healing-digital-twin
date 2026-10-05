"""Build ORAN_SPlane_PRISM_Review_v8 from the presented v7_FINAL.
Every new claim is sourced in the speaker notes; numbers come from project files verified 2026-10-03."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

NAVY, TEAL, KICK, GREY = "12263A", "18B89A", "1C7293", "5B6B7C"
BG, LIGHT, RULE = "F8FAFC", "EAF2F8", "D7E1EA"
AMBER, AMBERBG, RED, REDBG, GREEN, GREENBG, WHITE = "C58B18", "FFF8E8", "B3402F", "FBECE9", "1E8E5A", "E6F4EC", "FFFFFF"

def rgb(h): return RGBColor.from_string(h)

# ---------------------------------------------------------------- corrections
def replace_in_runs(slide, old, new):
    hits = 0
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if old in full and p.runs:
                p.runs[0].text = full.replace(old, new)
                for r in p.runs[1:]: r.text = ""
                hits += 1
    return hits

def replace_in_notes(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    hits = 0
    for p in tf.paragraphs:
        full = "".join(r.text for r in p.runs)
        if old in full and p.runs:
            p.runs[0].text = full.replace(old, new); [setattr(r, "text", "") for r in p.runs[1:]]; hits += 1
    return hits

def append_note(slide, text):
    tf = slide.notes_slide.notes_text_frame
    p = tf.add_paragraph(); p.text = text

S = list(prs.slides)  # v7 slides, 1-based index = v7 number
corr = [
 (23, "and correct fault attribution on seven of eight attack scenarios.", "and the exact fault named on 84 of 96 attack runs."),
 (4,  "End-application absolute time error, O-RU to primary reference", "Absolute time error at reference point E, common time standard"),
 (10, "tcpdump on both bridges — every PTP frame recorded unmodified", "tcpdump on both bridges during each run (raw PCAPs not kept in the packaged archive)"),
 (10, "pmc management telemetry per run", "pmc management queries per run (returned no data; see KPI sources)"),
 (11, "All namespaces share the host clock, so clocks cannot drift apart.", "All namespaces share the host clock, and free_running 1 is set: offset is measured, never steered."),
 (12, "tcpdump records every frame on brUP and brDN, unmodified", "tcpdump captures PTP frames on brUP and brDN; raw PCAPs not kept in the archive"),
 (13, "every PTP frame on brUP and brDN, unmodified", "brUP and brDN; raw PCAPs not kept in archive"),
 (13, "Partially available", "Configured, returned nothing"),
 (13, "management data set snapshots during the run", "GET requests logged; no responses returned in any run"),
 (13, "pmc management telemetry is absent for 48 of the 168 runs (all B3, C1, C2 and C3 replicates); this is recorded rather than imputed.",
      "pmc queries returned no data in any of the 168 runs."),
 (15, "and a superior priority2 value competes in the election.", "and a superior data set (randomised priority2 and clockClass) competes in the G.8275.1 alternate BMCA."),
 (19, "Recommended response labels only. No corrective action was executed on this testbed.", "Recommended response labels only. No corrective action was executed in this 168-run campaign."),
 (20, "Rule v3 was written after v2 false positives were observed, then re-hashed before scoring. It is reported as a transparent post-observation revision, not as independent validation.",
      "Rule v3 fixed v2 false positives seen on upstream captures and was re-hashed before scoring. v2, frozen before the campaign began, gives identical verdicts on all 168 runs."),
 (27, "The provisioned context holds a single expected boundary-clock identity, so a legitimate replacement relay is indistinguishable from an unauthorised one.",
      "The replacement identity is provisioned, but the base rule's rogue-BC clause reads only the primary BC identity, and additive-only v3 cannot overturn it."),
 (27, "48 of 168 runs lack pmc management telemetry.", "pmc management queries returned no data in any run."),
 (28, "No corrective action was executed or measured. ISOLATE, HOLDOVER and ESCALATE are verdict outputs, not actions taken.",
      "None in the 168-run campaign; a 13 Sep pilot ran a triggered port failover (5/5 vs 0/5)."),
 (25, "The machine-learning arm underperformed because of what this testbed can observe, not because the method is unsuited to the problem. Three measured constraints account for the result.",
      "The machine-learning arm underperformed on what this testbed can observe. Three measured constraints bear on the result; an equal-input comparison was not run, so method suitability itself is untested."),
 (29, "Provisioned replacement context", "Rogue-BC clause fix"),
 (29, "Extend the operator context to carry approved replacement boundary-clock identities, then re-freeze and re-run. Closes the 0/12 benign failure and is expected to raise specificity toward 0.979.",
      "Make the base rule's rogue-BC clause read the already-provisioned replacement identity, then re-freeze and re-run. Expected to raise specificity toward 0.979."),
]
log = []
for n, o, nw in corr:
    h = replace_in_runs(S[n-1], o, nw)
    log.append((n, h, o[:60]))
    if h == 0: print("CORRECTION NOT APPLIED:", n, o[:70])
replace_in_notes(S[15-1], "A3 replay has no direct counterpart in that set, which is stated rather than forced.",
                 "A3 replay has no counterpart in that ETSI set; replay is listed as a time-protocol threat in IETF RFC 7384 section 3.2.3.")
append_note(S[11-1], "v8 correction: every daemon config sets free_running 1 (cfg/g87251.base in g87251_testbed_v2.tgz), so ptp4l reports offset but never adjusts the clock; servo state stays s0 in every RU log.")
append_note(S[13-1], "v8 correction: all 40,320 pmc.jsonl lines across the 120 runs that have the file are outgoing 'sending: GET ...' requests with no response. Probable cause (not re-run): pmc was invoked without -d 24 while ptp4l ran on domain 24.")
append_note(S[20-1], "v8: FROZEN_V2.json frozen_at 2026-09-20T11:41:31Z; campaign_v3.log starts 11:43:15Z; EVALUATION_V4.json shows v2 macro sensitivity 0.9896, identical to v3. v2 was developed after earlier same-day live gap captures of the same C-scenario designs.")
append_note(S[27-1], "v8 correction: B_bc_replacement context.json carries expected_bc_identity_secondary = 020000fffe0000b1 and decision_rule_v3 adds it to the known set for D1. All 12 false ATTACK verdicts come from the base rule's 'Unauthorised clock inserted in the timing path' clause (A8), whose known_sources set reads only expected_bc_identity.")
append_note(S[28-1], "v8 correction: outputs/empirical_software_network_pilot_v1/S11_S12_CURRENT_REPORT.md records a detector-triggered port failover with matched no-action control (5/5 vs 0/5), on a separate software testbed with a jitter impairment. Its independent validation S15 was NOT established (no-trigger rate 8/15 below the frozen 0.8 criterion).")

# ---------------------------------------------------------------- primitives
LAYOUT = prs.slide_layouts[0]

def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = 0.08
    return s

def text(sl, x, y, w, h, runs, size=12, color=GREY, bold=False, font="Calibri", align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=0):
    """runs: str, or list of paragraphs; each paragraph str or list of (text, overrides) tuples."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space: p.space_after = Pt(space)
        segs = para if isinstance(para, list) else [(para, {})]
        for seg in segs:
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run(); r.text = t
            f = r.font; f.name = o.get("font", font); f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold); f.italic = o.get("italic", False); f.color.rgb = rgb(o.get("color", color))
    return tb

def frame(kicker, title, cite):
    sl = prs.slides.add_slide(LAYOUT)
    box(sl, 0, 0, 13.333, 7.5, fill=BG)
    box(sl, 0, 0, 13.333, 0.12, fill=TEAL)
    text(sl, 0.62, 0.23, 12.08, 0.25, kicker, size=12.8, color=KICK, bold=True)
    text(sl, 0.62, 0.49, 12.08, 0.6, title, size=33, color=NAVY, bold=True, font="Cambria")
    box(sl, 0.62, 1.16, 12.08, 0.48, fill=LIGHT, line=RULE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, 0.79, 1.2, 11.75, 0.42, cite, size=10.2, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    return sl

def table(sl, cols, headers, rows, y=1.82, row_h=None, size=11, first_font="Cambria", first_size=13.5,
          colors=None, header_size=11.5, gap=0.07, rx=0.62, rw=12.08):
    """cols: list of (x, w). rows: list of cell lists. Returns bottom y."""
    for (x, w), hd in zip(cols, headers):
        text(sl, x, y, w, 0.25, hd, size=header_size, color=KICK, bold=True)
    yy = y + 0.3
    box(sl, rx, yy, rw, 0.012, fill=RULE)
    yy += gap
    for ri, row in enumerate(rows):
        h = row_h[ri] if isinstance(row_h, list) else row_h
        for ci, ((x, w), cell) in enumerate(zip(cols, row)):
            if ci == 0:
                text(sl, x, yy, w, h, cell, size=first_size, color=NAVY, bold=True, font=first_font)
            else:
                c = (colors[ci] if colors else (NAVY if ci == len(cols) - 1 else GREY))
                text(sl, x, yy, w, h, cell, size=size, color=c)
        yy += h
        box(sl, rx, yy, rw, 0.012, fill=RULE)
        yy += gap
    return yy

def band(sl, y, h, label, body, dark=True, size=13):
    if dark:
        box(sl, 0.62, y, 12.08, h, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 0.85, y + 0.08, 2.3, h - 0.16, label, size=12, color=TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 3.1, y + 0.06, 9.4, h - 0.12, body, size=size, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    else:
        box(sl, 0.62, y, 12.08, h, fill=AMBERBG, line=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 0.85, y + 0.08, 2.3, h - 0.16, label, size=14, color=AMBER, bold=True, font="Cambria", anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 3.1, y + 0.06, 9.4, h - 0.12, body, size=size - 0.8, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

_NOTE_BODY = None
for _s in S:
    if _s.has_notes_slide and _s.notes_slide.notes_text_frame is not None:
        _NOTE_BODY = _s.notes_slide.notes_placeholder._element; break

def notes(sl, lines):
    ns = sl.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(copy.deepcopy(_NOTE_BODY))
    tf = ns.notes_text_frame
    tf.text = lines[0]
    for l in lines[1:]:
        tf.add_paragraph().text = l

NEW = {}

# ================================================================ R0 review map
sl = frame("REVIEW RESPONSE", "Reviewer Questions and Where Each Is Answered",
           "Questions transcribed from the reviewer's handwritten notes on the PRISM review presentation  |  Every new slide carries numbered references [n] listed at the end of the deck")
NEW["R0"] = sl  # rows filled after ordering (slide numbers)

# ================================================================ N1 split layers
sl = frame("SYSTEM ARCHITECTURE · PROTOCOL LAYERS", "O-RAN Functional Split: What the O-DU and O-RU Host",
           "ETSI TS 103 982 V8.0.0 (2024-01), O-RAN Architecture Description, clauses 3.1, 6.3.3–6.3.6, 6.4.7 [4]  |  3GPP TS 38.401 NG-RAN architecture [11]")
nodes = [
  ("O-CU-CP", "central unit · control plane", ["RRC", "PDCP (control plane)"]),
  ("O-CU-UP", "central unit · user plane", ["SDAP", "PDCP (user plane)"]),
  ("O-DU", "distributed unit", ["RLC", "MAC", "High-PHY"]),
  ("O-RU", "radio unit · physical node", ["Low-PHY  (FFT / iFFT, PRACH extraction)", "RF processing"]),
]
xs = [0.75, 3.9, 7.05, 10.2]
for (nm, sub, layers), x in zip(nodes, xs):
    hl = nm == "O-RU"
    box(sl, x, 1.95, 2.5, 3.0, fill=WHITE, line=(TEAL if hl else RULE), shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=(2.25 if hl else 1))
    text(sl, x + 0.15, 2.05, 2.2, 0.35, nm, size=19, color=NAVY, bold=True, font="Cambria")
    text(sl, x + 0.15, 2.42, 2.2, 0.25, sub, size=10.5, color=GREY)
    yy = 2.8
    for L in layers:
        box(sl, x + 0.15, yy, 2.2, 0.5, fill=(NAVY if hl else LIGHT), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, x + 0.22, yy, 2.05, 0.5, L, size=12, color=(WHITE if hl else NAVY), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        yy += 0.6
links = [("E1", 3.25), ("F1", 6.4), ("Open FH\n7-2x", 9.55)]
for lab, x in links:
    box(sl, x, 3.35, 0.65, 0.04, fill=TEAL)
    text(sl, x - 0.05, 3.45, 0.75, 0.5, lab, size=9.5, color=KICK, bold=True, align=PP_ALIGN.CENTER)
text(sl, 0.75, 5.08, 6.0, 0.3, "OPEN FRONTHAUL PLANES  (TS 103 982 cl. 6.4.7)", size=11.5, color=KICK, bold=True)
planes = [("C", "Control"), ("U", "User"), ("S", "Synchronization — this project"), ("M", "Management")]
px = 0.75
for k, v in planes:
    w = {"C": 1.55, "U": 1.45, "S": 3.3, "M": 2.15}[k]
    box(sl, px, 5.4, w, 0.42, fill=(TEAL if k == "S" else LIGHT), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, px + 0.1, 5.4, w - 0.2, 0.42, f"{k}-Plane · {v}", size=11.5, color=(WHITE if k == "S" else NAVY), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    px += w + 0.15
text(sl, 9.95, 5.08, 2.75, 0.8, "E2 terminates at O-CU-CP, O-CU-UP and O-DU, not at the O-RU (cl. 6.3.3–6.3.6).", size=10.5, color=GREY)
band(sl, 6.08, 0.95, "WHY THE O-RU",
     "Radio requirements are measured at the base-station antenna [8][9]; in O-RAN the antenna side is the O-RU (Low-PHY + RF). The S-plane is how the O-RU receives time, so it is where this project works.")
notes(sl, [
 "Answers reviewer question: explain the DU layers and the architecture.",
 "ETSI TS 103 982 V8.0.0 clause 3.1 definitions (verbatim): O-CU-CP hosts the RRC and the control plane part of the PDCP protocol; O-CU-UP hosts the user plane part of PDCP and SDAP; O-DU hosts RLC/MAC/High-PHY layers based on a lower layer functional split; O-RU hosts Low-PHY layer and RF processing (FFT/iFFT, PRACH extraction).",
 "Clause 6.3.6: the O-RU terminates the Open Fronthaul interface (LLS) and is a physical node. Clause 6.4.7: the Open FH interface includes the CUS planes and the M-plane. NOTE in 6.1: the LLS is split option 7-2x.",
 "Clauses 6.3.3, 6.3.4, 6.3.5: O-CU-CP, O-CU-UP and O-DU each terminate the E2 interface to the Near-RT RIC. Clause 6.3.6 (O-RU) lists only the Open Fronthaul interface, Low-PHY functions and the Open Fronthaul M-Plane; no E2.",
 "3GPP TS 38.133 (via Ruffini, ATIS 2018, slide 6): cell phase synchronization accuracy measured at BS antenna connectors shall be better than 3 us. 3GPP TS 38.104 Rel-19 clause 9.6.3.2: TAE limits 65 ns (MIMO), 260 ns (intra-band contiguous CA), 3 us (non-contiguous / inter-band CA)."])
NEW["N1"] = sl

# ================================================================ N2 LLS-C
sl = frame("SYSTEM ARCHITECTURE · TIMING DISTRIBUTION", "How the O-RU Receives Time: LLS-C1 to LLS-C4",
           "O-RAN.WG4.CUS.0 v06.00 synchronization configurations, as quoted in G. Armstrong, \"O-RAN Fundamentals\", ATIS WSTS 2023 [7]  |  Exposure column reasoned from [1], [2], [10], [12]")
cols = [(0.73, 1.35), (2.2, 5.6), (8.0, 4.6)]
rows = [
 ["LLS-C1", "O-DU is part of the synchronization chain. Network timing is distributed from O-DU to O-RU via a direct connection between O-DU site and O-RU site.",
  "A compromised O-DU or the O-DU–O-RU link controls the O-RU's time (T-FRHAUL-01 [2])."],
 ["LLS-C2", "O-DU is part of the synchronization chain. Timing goes from O-DU to O-RU; one or more Ethernet switches are allowed in the fronthaul network.",
  "As C1, plus any device on the switched fronthaul can inject or intercept PTP [10, sec. 4]."],
 ["LLS-C3", "O-DU is not part of the synchronization chain. Timing is distributed from PRTC/T-GM to O-RU, typically from central or aggregation sites.",
  "Grandmaster, boundary clocks and switches on the path are all exposed. TIMESAFE's production test used LLS-C3 [10]."],
 ["LLS-C4", "Synchronization reference provided to the O-RU with no involvement of the transport network, typically a local GNSS receiver.",
  "No PTP exposure; GNSS jamming or spoofing instead (RFC 7384 sec. 3.2.10 [1]; BSI [12])."],
]
yb = table(sl, cols, ["CONFIG", "DEFINITION (O-RAN.WG4.CUS.0, AS QUOTED IN [7])", "WHERE AN ATTACKER CAN ACT"], rows, row_h=0.84, size=12.5, first_size=16)
band(sl, yb + 0.3, 0.95, "OUR TESTBED",
     "Grandmaster → boundary clock → three radio units over bridged segments, with no O-DU in the timing chain: the LLS-C3 structure. The O-DU itself is not emulated.")
notes(sl, [
 "Answers reviewer questions: explain the architecture; which module will these attacks affect.",
 "Definitions are quoted from G. Armstrong, O-RAN Fundamentals, ATIS WSTS 2023, page 15; the specification O-RAN.WG4.CUS.0-v06.00 is cited on page 17. The O-RAN specification itself was not read directly; this is a secondary source and is labelled as such.",
 "The exposure column is our reasoning from the cited threat sources, not text from the O-RAN specification.",
 "Testbed topology: GM-A/GM-B on brUP, boundary clock BC between brUP and brDN, RU1-RU3 on brDN (topology.sh, slide on testbed architecture)."])
NEW["N2"] = sl

# ================================================================ N3 threat model
sl = frame("THREAT MODEL · WHO ATTACKS", "Who Attacks the S-Plane: Agents and Classes",
           "ETSI TR 104 106 V3.0.0 (2025-06) clause 7.2 threat agents and clause 7.4.1.2 [2]  |  IETF RFC 7384 (2014) clause 3.1 threat model [1]  |  TIMESAFE, ACM TOPS 2025, section 4 [10]")
text(sl, 0.73, 1.82, 5.6, 0.25, "THREAT AGENTS  (ETSI TR 104 106, CL. 7.2)", size=11.5, color=KICK, bold=True)
agents = [("Insiders", "malicious attack by a person with authorized access"),
          ("Nation-state", "aggressively target and gain persistent access to networks"),
          ("Cyber-criminals", "individuals who commit cybercrimes"),
          ("Hacktivists", "cyber-attacks for political or social gains"),
          ("Cyber-terrorists", "target O-RAN infrastructure through violence"),
          ("Script kiddies", "actors without deep technical expertise")]
yy = 2.15
for a, d in agents:
    text(sl, 0.73, yy, 1.85, 0.42, a, size=13.5, color=NAVY, bold=True, font="Cambria")
    text(sl, 2.6, yy + 0.03, 3.75, 0.42, d, size=11.5, color=GREY)
    yy += 0.47
box(sl, 6.55, 1.82, 0.012, 4.0, fill=RULE)
text(sl, 6.8, 1.82, 5.9, 0.25, "ATTACKER CLASSES BY ACCESS AND POSITION  (RFC 7384, CL. 3.1)", size=11.5, color=KICK, bold=True)
cls = [("Internal", "has access to a trusted segment of the network, or holds the keys"),
       ("External", "no keys; sees only encrypted or authenticated traffic"),
       ("Man-in-the-middle", "positioned to intercept and modify in-flight packets"),
       ("Injector", "not in-path; attacks by generating protocol packets")]
yy = 2.15
for a, d in cls:
    box(sl, 6.8, yy, 5.9, 0.5, fill=WHITE, line=RULE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, 6.95, yy, 1.95, 0.5, a, size=13, color=NAVY, bold=True, font="Cambria", anchor=MSO_ANCHOR.MIDDLE)
    text(sl, 8.9, yy, 3.7, 0.5, d, size=11, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.58
text(sl, 6.8, 4.5, 5.9, 0.25, "HOW THEY GET ON THE FRONTHAUL  ([10] sec. 4; [2] T-FRHAUL-02)", size=11.5, color=KICK, bold=True)
text(sl, 6.8, 4.8, 5.9, 1.05, [
   "• Remote access widened by multi-vendor, disaggregated equipment",
   "• Physical access to edge cell sites hosting O-RU / O-DU",
   "• Supply-chain backdoors in DU or RU equipment",
   "• Unauthorized access to the Open Fronthaul Ethernet interface"], size=11.5, color=NAVY)
band(sl, 6.0, 1.05, "ANSWER: WHO",
     "The standards name attackers by capability and position, not by identity. Six of our eight attacks fall in RFC 7384 rows that require an internal attacker; interception (C1) can also be done by an external man-in-the-middle.", size=12.2)
notes(sl, [
 "Answers reviewer questions: what kind of attack and by whom; who does these attacks; source of attack.",
 "ETSI TR 104 106 V3.0.0 clause 7.2 'Threat agent' lists: cyber-criminals, insiders, hacktivists, cyber-terrorists, script kiddies, nation-state (definitions quoted in short form).",
 "RFC 7384 clause 3.1.1: internal attackers have access to a trusted segment of the network or possess the encryption or authentication keys; external attackers do not have the keys and see only encrypted or authenticated traffic. Clause 3.1.2: MITM vs traffic injector.",
 "RFC 7384 Table 1: spoofing, replay, rogue master and time-protocol DoS are marked for internal MITM/injector only; interception/removal is marked for internal and external MITM. Mapping of our scenarios: A1, A8 rogue master; A2, C3 spoofing; A3 replay; A5 time-protocol DoS (all internal); C1 interception/removal (internal or external MITM); C2 malformed input falls under clause 3.2.11 'exploiting vulnerabilities', which Table 1 does not classify.",
 "TIMESAFE section 4: 'Once an attacker gains access to a device on the switched network, they can observe, replay, or inject PTP broadcast packets sent through the switch.' Pathways: multi-vendor remote access, physical proximity to edge gNBs, supply-chain backdoors in DU/RU.",
 "ETSI TR 104 106 T-FRHAUL-02: unauthorized access to Open Front Haul Ethernet L1 physical layer interface(s), threatening U-, S-, C- and M-plane traffic.",
 "We do not attribute any attack to a named real-world actor; no public source used here does."])
NEW["N3"] = sl

# ================================================================ N4 entry points
sl = frame("THREAT MODEL · WHERE AND HOW", "Attack Entry Points on the Testbed Timing Path",
           "Positions and tools read from the campaign scripts run/scenarios.sh and run/run_gap.sh in splane_campaign_CORRECTED_2026-09-20.tgz  |  Attacker classes per RFC 7384 cl. 3.1 [1]")
# diagram (left)
def node(x, y, w, h, t, sub="", dark=False):
    box(sl, x, y, w, h, fill=(NAVY if dark else LIGHT), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, x, y + 0.04, w, 0.3, t, size=14, color=(WHITE if dark else NAVY), bold=True, font="Cambria", align=PP_ALIGN.CENTER)
    if sub: text(sl, x, y + 0.33, w, 0.25, sub, size=9.5, color=(WHITE if dark else GREY), align=PP_ALIGN.CENTER)
node(0.75, 2.05, 1.35, 0.65, "GM-A", "grandmaster")
node(0.75, 3.0, 1.35, 0.65, "GM-B", "backup GM")
box(sl, 2.3, 1.95, 0.22, 2.65, fill=TEAL); text(sl, 2.16, 4.62, 0.5, 0.25, "brUP", size=9.5, color=KICK, bold=True, align=PP_ALIGN.CENTER)
node(2.75, 2.85, 1.35, 0.65, "BC", "boundary clock", dark=True)
box(sl, 4.3, 1.95, 0.22, 2.65, fill=TEAL); text(sl, 4.16, 4.62, 0.5, 0.25, "brDN", size=9.5, color=KICK, bold=True, align=PP_ALIGN.CENTER)
for i, y in enumerate([1.95, 2.85, 3.75]):
    node(4.75, y, 1.3, 0.65, f"RU{i+1}", "radio unit")
def marker(x, y, w, n, t):
    box(sl, x, y, w, 0.55, fill=REDBG, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5)
    text(sl, x + 0.06, y, 0.3, 0.55, n, size=16, color=RED, bold=True, font="Cambria", anchor=MSO_ANCHOR.MIDDLE)
    text(sl, x + 0.36, y, w - 0.42, 0.55, t, size=10, color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
marker(2.62, 1.98, 1.6, "3", "In-path at BC port")
marker(2.62, 3.72, 1.6, "2", "Device across both bridges")
marker(4.4, 4.95, 1.75, "1", "Injector on RU segment")
# table (right)
cols = [(6.45, 0.45), (6.95, 1.5), (8.5, 4.2)]
rows = [
 ["1", "A1, A2, A3, A5, C2, C3", "Device on the RU-side fronthaul segment (the ru3 namespace, or an added namespace on brDN). A1: second ptp4l with off-allow-list identity; A2, A5, C2, C3: Scapy-built frames; A3: tcpdump capture, then tcpreplay."],
 ["2", "A8", "Two-port ptp4l bridging brUP and brDN that slaves to the real GM and re-advertises downstream with stepsRemoved incremented."],
 ["3", "C1", "Boundary clock's downstream port set down for 60–75 % of the injection window (as drawn): every frame is removed, not just selected message types."],
]
yb = table(sl, cols, ["#", "SCENARIOS", "HOW IT WAS LAUNCHED"], rows, y=1.82, row_h=[1.25, 0.82, 0.82], size=11, first_size=15, rx=6.45, rw=6.25)
text(sl, 6.45, yb + 0.02, 6.25, 0.55, "RFC 7384 position: 1 = internal injector · 2 = internal man-in-the-middle (rogue master) · 3 = man-in-the-middle (internal or external)", size=10.5, color=KICK, bold=True)
band(sl, 6.05, 1.0, "SCOPE",
     "All attackers are Linux network namespaces on one host; no real network or host clock was touched. Real-world access paths are those on the threat-model slide.", dark=False, size=13)
notes(sl, [
 "Answers reviewer questions: who does these attacks and how; source of attack.",
 "Read from scenarios.sh (A1 rogue ptp4l in netns 'rogue' attached to brDN; A2 inject.py from ru3; A3 tcpdump on brDN then tcpreplay from ru3; A5 flood.py from ru3; A8 two-port ptp4l in netns 'rbc' attached to both bridges) and run_gap.sh (C1 'ip link set v-bc-dn down' for c1_down_pct of the window; C2 inject_malformed.py and C3 inject_wholesecond.py from ru3). inject.py, flood.py, inject_malformed.py and inject_wholesecond.py import scapy.all (Ether, Raw, sendp).",
 "C1 limitation stated on the slide: the blackhole removes all frames on the port, so it is a coarse realisation of ETSI T-SPLANE-04 'selective interception and removal'. randparams.py offers 55-75 %; the 168 drawn replicates used 60-75 %.",
 "Clock safety: run_gap.sh header states netns only, software timestamping, host clock never touched."])
NEW["N4"] = sl

# ================================================================ N5 provenance
sl = frame("THREAT MODEL · PROVENANCE", "Where Each of the Eight Attacks Comes From",
           "IETF RFC 7384 cl. 3.2 [1]  |  ETSI TR 104 106 V3.0.0 T-SPLANE-01…04 [2]  |  ETSI TS 104 105 V7.0.0 O-RAN Security Test Specifications [3]  |  TIMESAFE sec. 4.1 [10]")
cols = [(0.73, 0.55), (1.32, 2.15), (3.55, 3.6), (7.25, 2.7), (10.05, 2.65)]
rows = [
 ["A1", "Rogue grandmaster", "RFC 7384 §3.2.4 rogue master; T-SPLANE-03", "TS 104 105 §11.1.5.2.2 Rogue PTP Instance; TIMESAFE spoofing", "Randomised identity, priority2, clockClass"],
 ["A2", "Sync / Follow_Up spoofing", "RFC 7384 §3.2.2 spoofing", "No ETSI test of this exact form", "Off-allow-list source, bursts, no Announce"],
 ["A3", "Replay", "RFC 7384 §3.2.3 replay", "TIMESAFE replay attack", "Live capture replayed with tcpreplay"],
 ["A5", "DoS / PTP flooding", "RFC 7384 §3.2.9 time-protocol DoS; T-SPLANE-01", "TS 104 105 §11.1.5.1.1, §24.2.1.1", "Burst size and gap randomised"],
 ["A8", "Rogue boundary clock", "RFC 7384 §3.2.4 rogue master, in MITM position", "Adapted: no TS 104 105 test places the rogue as a relay", "Relays the real GM, stepsRemoved +1"],
 ["C1", "Interception and removal", "RFC 7384 §3.2.5; T-SPLANE-04", "TS 104 105 §11.1.5.3.1", "Realised as a port blackhole (all frames)"],
 ["C2", "Malformed frames", "RFC 7384 §3.2.11 exploiting vulnerabilities", "TS 104 105 §24.2.1.2 S-Plane PTP Unexpected Input", "versionPTP 3, reserved type, short length"],
 ["C3", "Whole-second field abuse", "RFC 7384 §3.2.2 spoofing (false time)", "None found; fields per IEEE 1588-2019 §8.2.4", "Our construction: leap61 set, UTC offset 0"],
]
yb = table(sl, cols, ["ID", "ATTACK", "PUBLISHED THREAT", "PUBLISHED TEST OR STUDY", "WHAT WE CHOSE"], rows, row_h=0.47, size=10.5, first_size=13.5, gap=0.035,
           colors=[None, NAVY, GREY, GREY, NAVY])
band(sl, 6.38, 0.82, "ANSWER",
     "None of the eight threat classes was invented. Seven follow a published threat or test; C3 is our own instance of the RFC 7384 spoofing class, with no published test of that exact form.", size=12.4)
notes(sl, [
 "Answers reviewer questions: were the 8 attacks already existing or made up; where were they referred from; citation of the attack references.",
 "RFC 7384 section numbers verified against rfc-editor.org text: 3.2.1 packet manipulation, 3.2.2 spoofing, 3.2.3 replay, 3.2.4 rogue master, 3.2.5 packet interception and removal, 3.2.6 packet delay manipulation, 3.2.7 L2/L3 DoS, 3.2.8 cryptographic performance, 3.2.9 DoS against the time protocol, 3.2.10 grandmaster time source attack, 3.2.11 exploiting vulnerabilities, 3.2.12 network reconnaissance.",
 "ETSI TR 104 106 V3.0.0 clause 7.4.1.2: T-SPLANE-01 DoS against a Master clock; T-SPLANE-02 impersonation of a Master clock; T-SPLANE-03 rogue PTP instance wanting to be Grand Master; T-SPLANE-04 selective interception and removal of PTP timing packets.",
 "ETSI TS 104 105 V7.0.0: 11.1.5.1.1 DOS Master Clock LLS C1 C2 C3; 11.1.5.2.1 Impersonation Master Clock; 11.1.5.2.2 Rogue PTP Instance; 11.1.5.3.1 Selective Interception and Removal of PTP Timing Packets; 11.1.5.3.2 Delay Attack on PTP Timing Packets; 24.2.1.1 S-Plane PTP DoS Attack; 24.2.1.2 S-Plane PTP Unexpected Input.",
 "TIMESAFE section 4.1 implements spoofing (fake Announce with manipulated BMCA attributes) and replay (sniffed Sync/Follow_Up retransmitted later).",
 "Catalogue classes not run: A4 delay attack = RFC 7384 3.2.6 and TS 104 105 11.1.5.3.2; A6/A7 GNSS spoofing/jamming = RFC 7384 3.2.10.",
 "C3 statement is bounded: we did not find a published test of this exact field combination in the sources listed; that is not proof none exists."])
NEW["N5"] = sl

# ================================================================ N6 impact
sl = frame("THREAT MODEL · IMPACT", "Which Component Each Attack Affects, and What Happens",
           "Measured: ptp4l logs of RU1–RU3, 12 replicates per scenario (36 RU logs each), extracted_168run_ptp4l_logs  |  Impact classes: RFC 7384 Table 1 [1]; ETSI T-SPLANE impacts [2]; TIMESAFE [10]")
cols = [(0.73, 0.55), (1.32, 2.6), (4.0, 1.75), (5.85, 3.55), (9.5, 3.2)]
rows = [
 ["A1", "All RU slave clocks on the segment", "False time", "36/36 RU logs selected the rogue as best master", "TIMESAFE: O-RU crash ≈2 s after spoofing, cell drop, manual reboot"],
 ["A8", "All RUs behind the inserted clock", "False time", "36/36 RU logs took the rogue BC as a new master", "T-SPLANE-03: slaves given inaccurate time"],
 ["C1", "RUs downstream of the BC port", "Accuracy loss, DoS", "36/36 RU ports lost master on Announce timeout; servo reports 7.8 vs 22 per log", "T-SPLANE-04: degradation or free-running"],
 ["C3", "RU time-properties (UTC offset)", "False time", "No master change; servo reports 3.0 vs 22 per log", "Spoofing class: false time"],
 ["A2 A3 A5 C2", "RU receive path", "False time / DoS", "No master change in any RU log; visible only in packet evidence", "A5: T-SPLANE-01 clock service interrupted; A3: TIMESAFE replay"],
]
yb = table(sl, cols, ["ID", "AFFECTED COMPONENT", "IMPACT CLASS", "MEASURED IN OUR RUNS", "REPORTED IN LITERATURE"], rows, row_h=[0.6, 0.5, 0.6, 0.5, 0.6], size=10.8, first_size=13.5, gap=0.04,
           colors=[None, NAVY, GREY, NAVY, GREY])
box(sl, 0.62, yb + 0.05, 12.08, 0.55, fill=WHITE, line=RULE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(sl, 0.8, yb + 0.05, 11.8, 0.55, [[("Service limits at stake:  ", {"bold": True, "color": NAVY}),
     ("TDD cell phase sync better than 3 µs (3GPP TS 38.133 [8]) · TAE ≤ 65 ns for MIMO, ≤ 260 ns intra-band contiguous CA (3GPP TS 38.104 §9.6.3.2 [9])", {})]],
     size=11, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
band(sl, 6.38, 0.82, "LIMIT",
     "Every daemon ran with free_running 1: offsets were measured, never steered. Measured impact is protocol-level; time-error and service impact come from the cited literature.", dark=False, size=13)
notes(sl, [
 "Answers reviewer questions: what are the impacts; which module will these attacks affect.",
 "Measured counts computed on 2026-10-03 from ml_comparison_input/extracted_168run_ptp4l_logs (RU1-RU3 logs, 12 replicates per scenario): A1 36/36 logs contain 'selected best master clock 020000.fffe.0000xx' with a non-provisioned identity; A8 36/36 contain 'new foreign master' with the rogue BC identity; C1 36/36 contain ANNOUNCE_RECEIPT_TIMEOUT_EXPIRES. Mean servo report lines per RU log: baseline 22.0, C1 7.8, C3 3.0.",
 "Benign comparison from the same logs: B3 congestion 12/36 RU logs hit the Announce timeout; B_bc_replacement 36/36 during the planned swap. A timeout alone therefore does not prove an attack.",
 "All daemon configs set free_running 1 (cfg/g87251.base); servo state is s0 in every RU log. No RU clock was steered, so false-time impact was not realised physically here.",
 "Literature (TIMESAFE, arXiv 2412.13049v3, verified verbatim 4 Oct): production network is O-RAN LLS-C3; with all ports PTP dynamic: 'Approximately 2 seconds after the attack begins, the RU crashes, stopping its operations, causing the 5G cell to drop and the UE to lose connection'; in a separate run 'around second 440 there is a 50% drop in throughput, progressing to a 75% drop by second 510, and ultimately causing the base station to crash at around second 580'; another run reports a 50% drop about 380 s after attack start. ETSI TR 104 106 impact texts quoted for T-SPLANE-01, -03, -04. RFC 7384 Table 1 impact columns: false time, accuracy degradation, DoS.",
 "Service limits: 3GPP TS 38.133 as quoted by Ruffini (ATIS 2018, slide 6); TS 38.104 Rel-19 clause 9.6.3.2."])
NEW["N6"] = sl

# ================================================================ N7 KPI sources
sl = frame("KPI SOURCES", "KPIs Are Measured From the Protocol, Not a Model",
           "IEEE 1588-2019 message formats and data sets [5]  |  linuxptp ptp4l and pmc [14]  |  Extractor ptp_deep_extract.py; logs extracted_168run_ptp4l_logs; pmc_log.sh")
srcs = [
 ("01", "Wire capture", "tcpdump on brUP / brDN → ptp_deep_extract.py", "56 fields per PTP message: header, flags, correctionField, sourcePortIdentity, sequenceId, Announce data set, G.8275.1 conformance flags", "Used by the rule (ARM A)", GREEN),
 ("02", "Daemon servo log", "ptp4l -m output on each node", "master offset (ns), servo state, frequency adjustment (ppb), path delay (ns), port-state transitions", "Used by the ML model (ARM B)", GREEN),
 ("03", "Management queries", "pmc GET on 7 data sets per node, every second", "Intended: CURRENT, PARENT, TIME_PROPERTIES, PORT data sets and port statistics", "Returned no data in any run", RED),
 ("04", "Operator context", "context.json written per run", "Provisioned allow-list, expected boundary clock, maintenance-window state, ground-truth label", "Configuration, not measurement", AMBER),
]
x = 0.73
for num, ttl, how, what, use, c in srcs:
    box(sl, x, 1.85, 2.9, 3.85, fill=WHITE, line=RULE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, x + 0.18, 1.95, 0.6, 0.35, num, size=17, color=KICK, bold=True, font="Cambria")
    text(sl, x + 0.18, 2.32, 2.6, 0.35, ttl, size=17, color=NAVY, bold=True, font="Cambria")
    text(sl, x + 0.18, 2.72, 2.6, 0.55, how, size=10.5, color=KICK, bold=True)
    text(sl, x + 0.18, 3.3, 2.6, 1.6, what, size=11.2, color=GREY)
    bg = {GREEN: GREENBG, RED: REDBG, AMBER: AMBERBG}[c]
    box(sl, x + 0.15, 5.05, 2.6, 0.48, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, x + 0.22, 5.05, 2.5, 0.48, use, size=10.5, color=c, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    x += 3.05
text(sl, 0.73, 5.82, 12.0, 0.45, [[("Example servo line, baseline run 1, RU1:  ", {"bold": True, "color": NAVY}),
     ("ptp4l[517.414]: master offset 1191 s0 freq -1547 path delay 5313", {"font": "Courier New", "color": NAVY})]], size=11, color=GREY)
band(sl, 6.38, 0.82, "ANSWER",
     "No model generates the KPIs. They are read from PTP itself, defined by IEEE 1588-2019 and ITU-T G.8275.1. The ML model consumes six servo-derived features; it does not produce KPIs.", size=12.4)
notes(sl, [
 "Answers reviewer question: from which model did we get the KPI.",
 "Deep extract columns (56) verified from the CSV header: capture_ts_ns, eth_src, eth_dst, vlan_id, vlan_pcp, message_type, message_type_raw, version_ptp, minor_version_ptp, message_length, domain_number, major/minor_sdo_id, flag_field_hex and 12 individual flags, correction_ns, source_clock_identity, source_port_number, sequence_id, control_field, log_message_interval, origin_ts_ns, current_utc_offset, gm_priority1, gm_clock_class, gm_clock_accuracy, gm_offset_scaled_log_variance, gm_priority2, gm_clock_identity, steps_removed, time_source, tlv_types, has_path_trace, path_trace_hops, requesting_clock_identity, requesting_port_number, clock_class_legal, bmca_fields_all_zero, gm_identity_equals_source, g87251_domain_ok, g87251_mcast_ok, g87251_priority1_ok, g87251_log_sync_ok, g87251_log_announce_ok, g87251_log_delayreq_ok.",
 "pmc: pmc_log.sh queries TIME_STATUS_NP, CURRENT_DATA_SET, PARENT_DATA_SET, TIME_PROPERTIES_DATA_SET, PORT_DATA_SET, PORT_STATS_NP and PORT_SERVICE_STATS_NP per node. All 40,320 logged lines are unanswered requests. Probable cause, not re-run: pmc invoked without -d 24 while ptp4l used domainNumber 24.",
 "ML features used (ML_VS_RULE_COMPARISON.md): offset_mean, offset_std, offset_abs_max, path_delay_mean, pdv_std, holdover_rate.",
 "RAN KPIs collected over E2 are not available at the S-plane: E2 does not terminate at the O-RU (ETSI TS 103 982 clause 6.1)."])
NEW["N7"] = sl

# ================================================================ N8 KPI classification
sl = frame("KPI CLASSIFICATION FOR RECOVERY", "KPI Classes and the Recovery Decision Each One Drives",
           "Classes derived from the clause order of run/decision_rule.py (base) and run/decision_rule_v3.py  |  Field legality per IEEE 1588-2019 [5] and ITU-T G.8275.1 [6]  |  Response labels as defined on the classification-method slide")
cols = [(0.73, 2.0), (2.8, 3.95), (6.85, 2.75), (9.7, 3.0)]
rows = [
 ["Protocol legality", "version_ptp, message_type_raw, message_length, clock_class_legal, g87251_domain_ok, g87251_mcast_ok, g87251_priority1_ok, steps_removed ≥ 255, flag_alternateMaster", "Structurally illegal traffic (C2; illegal Announce)", "ATTACK → ISOLATE source"],
 ["Identity and topology", "source_clock_identity, gm_clock_identity, gm_priority2, gm_clock_class, steps_removed vs provisioned inventory", "Rogue GM (A1), rogue BC (A8), unknown sender (A2)", "ATTACK → ISOLATE source"],
 ["Continuity and rate", "sequence_id regressions; delivered message rate vs the rate declared in log_message_interval", "Replay (A3), flooding (A5), removal (C1)", "ATTACK → ISOLATE source"],
 ["Time properties", "flag_leap61 / leap59, current_utc_offset, flag_currentUtcOffsetValid, flag_timeTraceable", "Whole-second abuse (C3)", "ATTACK → ISOLATE source"],
 ["Operator context", "context.json: gm_allowlist, expected_bc_identity, maintenance_window_open", "Planned vs unplanned change (B2, B7, B_unplanned)", "BENIGN → TOLERATE / HOLDOVER · UNKNOWN → ESCALATE"],
 ["Timing quality", "ptp4l servo log: master offset, path delay; derived offset_std, pdv_std", "Degradation magnitude; not used for the verdict", "Severity input only"],
 ["Not observable here", "GNSS status, SyncE quality level, oscillator error", "Needed for A6, A7, B1, B4, B6", "Requires hardware"],
]
yb = table(sl, cols, ["KPI CLASS", "KPIs (FIELD, LOG OR CONTEXT-KEY NAMES)", "WHAT IT REVEALS", "RECOVERY DECISION"], rows, row_h=[0.66, 0.5, 0.5, 0.5, 0.5, 0.45, 0.4], size=10.5, first_size=13, gap=0.035,
           colors=[None, GREY, GREY, NAVY])
band(sl, 6.4, 0.8, "READING ORDER",
     "Legality, identity, continuity and time-property KPIs decide ATTACK; context separates planned from unplanned changes; timing quality grades severity. Actions are recommended, not executed, in the campaign.", size=11.8)
notes(sl, [
 "Answers reviewer request: classify the KPIs so that we can proceed with recovery.",
 "Mapping to code: base rule ordered clauses (decision_rule.py): rate flood -> A5; sequenceId regressions -> A3; unauthorised relaying clock -> A8; unauthorised self-announcing GM -> A1; illegal clockClass, priority1 != 128, stepsRemoved >= 255, alternateMasterFlag -> ATTACK; off-allow-list GM -> A1; unknown sender -> A2; out-of-profile domain/multicast/cadence -> A5; GM change with maintenance -> BENIGN B2; GM change without maintenance -> UNKNOWN. v3 adds D1 rate starvation (C1), D2 malformed frames (C2), D3 whole-second abuse (C3).",
 "Timing-quality KPIs are not used by the rule's verdict; they are the ML arm's inputs. Recovery labels: ATTACK -> ISOLATE, BENIGN -> TOLERATE/HOLDOVER, UNKNOWN -> ESCALATE (classification-method slide). Only the 13 Sep pilot executed an action (port failover)."])
NEW["N8"] = sl

# ================================================================ N10 dataset provenance
sl = frame("DATASET PROVENANCE", "Where Every Dataset Came From and How It Is Used",
           "dataset/README.md (classification of 16 Sep)  |  TIMESAFE data repository github.com/genesys-neu/s-plane_security [10]  |  outputs/empirical_software_network_pilot_v1/START_HERE_FINAL.md")
cols = [(0.73, 2.35), (3.15, 4.05), (7.3, 3.05), (10.45, 2.25)]
rows = [
 ["168-run campaign", "Generated by this team on the software testbed (six linuxptp 4.0 daemons, G.8275.1), 20 Sep 2026. Archive sha256 6149b4fb…, 1,096 files, 1,451,909 packet records.", "Every result in this deck", "PRIMARY EVIDENCE"],
 ["Testbed configs", "cfg/*.cfg for all nodes in g87251_testbed_v2.tgz (17 Sep); not inside the corrected campaign archive.", "Configuration provenance", "REFERENCE"],
 ["TIMESAFE captures", "Public data of Groen et al., real hardware: Foxconn O-RU, NVIDIA Aerial DU, Dell S5248F-ON switch, Qulsar QG-2 grandmaster [10].", "August model work; profile cross-check", "REFERENCE ONLY"],
 ["13 Sep pilot", "Second software testbed: 25 S14 runs, 232,403 packets; S11 closed-loop trial; S15 validation not established.", "Closed-loop feasibility", "DEVELOPMENT DATA"],
 ["Simulator output", "Synthetic physics-simulator rows; quarantined in dataset/illegitimate.", "Software testing only", "NOT EVIDENCE"],
]
yb = table(sl, cols, ["DATASET", "ORIGIN", "USED FOR", "STATUS"], rows, row_h=[0.86, 0.66, 0.86, 0.66, 0.6], size=12, first_size=15, gap=0.05,
           colors=[None, GREY, NAVY, KICK])
band(sl, 6.38, 0.82, "ANSWER",
     "The evaluated dataset was generated by us on a reproducible testbed. The only external data, TIMESAFE, is public and cited; August results that used it were withdrawn on 17 Sep.", size=12.4)
notes(sl, [
 "Answers reviewer question: where did we get the dataset.",
 "Archive splane_campaign_CORRECTED_2026-09-20.tgz, sha256 6149b4fb15940cdac694ba8666d97041b4eb62d5523d2631c1edf96d531e4ed9 (verified on disk 2026-10-03). Packet-record total from the dataset slide.",
 "TIMESAFE repository clone at dataset/legitimate/timesafe_real_hardware_captures/s-plane_security_repo, git remote https://github.com/genesys-neu/s-plane_security.git; equipment per TIMESAFE section 5.1.",
 "Walkthrough of 17 Sep withdrew the August figures: some 'unseen' sessions duplicated training files and 12 of 28 features carried information only in simulated data.",
 "13 Sep pilot: S14 workbook 232,403 packets, 25 prospective runs; S15 independent validation not established (receiver-below-boundary no-trigger 8/15 vs required 0.8)."])
NEW["N10"] = sl

# ================================================================ N9 applicability
sl = frame("APPLICABILITY", "Does the Method Work Only for O-RAN?",
           "IEEE 1588 PTP profiles list, IEEE SA [13]  |  IETF RFC 7384 scope: time protocols, focusing on PTP and NTP [1]  |  Tested in this project: ITU-T G.8275.1 only [6]")
text(sl, 0.73, 1.82, 5.6, 0.25, "OTHER NETWORKS THAT RUN PTP  (IEEE 1588 PROFILES [13])", size=11.5, color=KICK, bold=True)
prof = [("Telecom", "ITU-T frequency profile; phase/time with full and partial timing support"),
        ("Power", "IEEE C37.238; IEC utility-automation profiles"),
        ("Broadcast and audio", "SMPTE broadcast profile; AES67 media profile"),
        ("A/V, industry, automotive", "IEEE Std 802.1AS; Avnu; AUTOSAR"),
        ("Enterprise, test", "IETF enterprise profile; LXI; GigE Vision")]
yy = 2.15
for a, d in prof:
    text(sl, 0.73, yy, 2.25, 0.5, a, size=13, color=NAVY, bold=True, font="Cambria")
    text(sl, 3.0, yy + 0.02, 3.3, 0.55, d, size=11, color=GREY)
    yy += 0.66
box(sl, 6.55, 1.82, 0.012, 4.3, fill=RULE)
blocks = [("TRANSFERS AS IS", GREEN, GREENBG, "IEEE 1588 message legality, sequenceId continuity, delivered vs declared rate, identity allow-list and maintenance context. The RFC 7384 threat list covers PTP generally."),
          ("MUST BE RE-DERIVED PER PROFILE", AMBER, AMBERBG, "G.8275.1 constants: domain 24–43, priority1 128, Sync 16/s, Announce 8/s, alternate BMCA, L2 multicast addressing."),
          ("NOT DEMONSTRATED", RED, REDBG, "Any profile other than G.8275.1; hardware timestamping; non-O-RAN deployments.")]
yy = 1.82
for t, c, bg, d in blocks:
    box(sl, 6.8, yy, 5.9, 1.3, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, 7.0, yy + 0.1, 5.5, 0.28, t, size=11.5, color=c, bold=True)
    text(sl, 7.0, yy + 0.42, 5.55, 0.85, d, size=11.5, color=NAVY)
    yy += 1.45
band(sl, 6.38, 0.82, "ANSWER",
     "The checks belong to PTP, not to O-RAN, so the method is portable by design after re-deriving profile constants. It has been demonstrated only on the G.8275.1 telecom profile.", size=12.4)
notes(sl, [
 "Answers reviewer question: does it only work on O-RAN or any other architecture.",
 "IEEE SA PTP profiles page lists: Default Delay Request-Response, Default Peer-to-Peer, High Accuracy Default (IEEE); ITU-T telecom profiles for frequency, phase/time with full timing support, time/phase with partial timing support; IETF Enterprise Profile; SMPTE broadcast profile; AES67 Media Profile; IEEE Std 802.1AS; IEEE C37.238 Power Profile; IEC Power Utility Automation and PRP/HSR automation profiles; IEC 'U' and 'D' industrial profiles; Avnu automotive; AUTOSAR; LXI; GigE Vision.",
 "RFC 7384 abstract: security requirements for time protocols, focusing on PTP and NTP.",
 "Portability is a design argument. The rule hard-codes G.8275.1 constants (domain, priority1, message intervals) and was run only on that profile."])
NEW["N9"] = sl

# ================================================================ references
refs = [
 "[1] T. Mizrahi, \"Security Requirements of Time Protocols in Packet Switched Networks,\" IETF RFC 7384 (Informational), Oct. 2014.",
 "[2] ETSI TR 104 106 V3.0.0 (2025-06), O-RAN Security Threat Modeling and Risk Assessment (O-RAN WG11).",
 "[3] ETSI TS 104 105 V7.0.0 (2025-06), O-RAN Security Test Specifications.",
 "[4] ETSI TS 103 982 V8.0.0 (2024-01), O-RAN Architecture Description.",
 "[5] IEEE Std 1588-2019, Precision Clock Synchronization Protocol for Networked Measurement and Control Systems.",
 "[6] ITU-T G.8275.1, Precision time protocol telecom profile for phase/time synchronization with full timing support from the network.",
 "[7] G. Armstrong, \"O-RAN Fundamentals,\" ATIS WSTS 2023, quoting O-RAN.WG4.CUS.0 v06.00.",
 "[8] S. Ruffini, \"Sync for 5G: What Is Needed,\" ATIS 2018, quoting 3GPP TS 38.133 cell phase synchronization accuracy.",
 "[9] 3GPP TS 38.104 Rel-19, NR Base Station radio transmission and reception, clause 9.6.3.2 Time alignment error.",
 "[10] J. Groen et al., \"TIMESAFE: Timing Interruption Monitoring and Security Assessment for Fronthaul Environments,\" ACM TOPS 28(5), 2025, doi:10.1145/3775060 (arXiv 2412.13049v3); data: github.com/genesys-neu/s-plane_security.",
 "[11] ETSI TS 138 401 (3GPP TS 38.401) Rel-18, NG-RAN Architecture description.",
 "[12] BSI (Federal Office for Information Security), 5G Risk Analysis: Radio Access Network, v1.0, 20 Feb 2025.",
 "[13] IEEE SA, IEEE 1588 PTP profiles list, sagroups.ieee.org/1588/ptp-profiles.",
 "[14] linuxptp project, ptp4l and pmc documentation, linuxptp.nwtime.org.",
]
sl = frame("REFERENCES", "References for the Review-Response Slides",
           "Numbered as cited on the new slides  |  Sources of the original slides remain on those slides  |  Verified on 3 October 2026")
text(sl, 0.73, 1.85, 5.85, 4.9, refs[:8], size=12.5, color=NAVY, space=10)
text(sl, 6.85, 1.85, 5.85, 4.9, refs[8:], size=12.5, color=NAVY, space=10)
notes(sl, ["URLs: RFC 7384 rfc-editor.org/rfc/rfc7384; ETSI TR 104 106 etsi.org/deliver/etsi_tr/104100_104199/104106/03.00.00_60/; ETSI TS 104 105 etsi.org/deliver/etsi_TS/104100_104199/104105/07.00.00_60/; ETSI TS 103 982 etsi.org/deliver/etsi_ts/103900_103999/103982/08.00.00_60/; Armstrong wsts.atis.org/wp-content/uploads/2023/02/Greg-Armstrong_ORAN-Fundamentals.pdf; Ruffini tam.atis.org/wp-content/uploads/2018/10/1_02_Ericsson_Ruffini_Sync-5G-What-Is-Needed.pdf; TIMESAFE arxiv.org/abs/2412.13049 and doi.org/10.1145/3775060; BSI 5G RAN risk analysis (local copy in literature-survey/papers/Government-Reports).",
           "IEEE 1588-2019 and ITU-T G.8275.1 are paywalled; clause-level claims about them come from the project's earlier verified material; the full texts were not re-read for v8."])
NEW["REF"] = sl

# ---------------------------------------------------------------- ordering
order = [("v",1),("v",2),("n","R0"),("v",3),("n","N1"),("n","N2"),("v",4),("v",5),("n","N3"),("n","N4"),("n","N5"),("n","N6"),
         ("v",6),("v",7),("v",8),("v",9),("v",10),("v",11),("v",12),("n","N7"),("n","N8"),("v",13),("n","N10"),
         ("v",14),("v",15),("v",16),("v",17),("v",18),("v",19),("v",20),("v",21),("v",22),("v",23),("v",24),("v",25),
         ("v",26),("v",27),("v",28),("v",29),("n","N9"),("v",30),("n","REF")]
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
v7_ids = ids[:30]
new_map = {k: ids[30 + i] for i, k in enumerate(["R0","N1","N2","N3","N4","N5","N6","N7","N8","N10","N9","REF"])}
for el in ids: sldIdLst.remove(el)
pos = {}
for i, (t, k) in enumerate(order, 1):
    sldIdLst.append(v7_ids[k-1] if t == "v" else new_map[k]); pos[(t, k)] = i

# fill R0 now that numbers are known
sl = NEW["R0"]
P = lambda *ks: ", ".join(str(pos[("n", k)]) for k in ks)
cols = [(0.73, 7.5), (8.35, 4.35)]
rows = [
 ["What kind of attack, by whom, and how? Who does these attacks? Source of attack?", f"Threat model and entry points: slides {P('N3','N4')}"],
 ["Were the 8 attacks existing or made up? Where were they referred from? Citations?", f"Attack provenance: slide {P('N5')}; references: slide {P('REF')}"],
 ["What are the impacts? Which module will these attacks affect?", f"Impact and affected component: slide {P('N6')}"],
 ["Explain the DU layers and the architecture", f"Functional split and timing configurations: slides {P('N1','N2')}"],
 ["Does it only work on O-RAN, or on other architectures?", f"Applicability: slide {P('N9')}"],
 ["From which model did we get the KPIs?", f"KPI sources: slide {P('N7')}"],
 ["Where did we get the dataset?", f"Dataset provenance: slide {P('N10')}"],
 ["Classify the KPIs so that we can proceed with recovery", f"KPI classification for recovery: slide {P('N8')}"],
]
yb = table(sl, cols, ["REVIEWER QUESTION", "ANSWERED ON"], rows, row_h=0.46, size=13, first_font="Calibri", first_size=13, gap=0.04,
           colors=[None, NAVY])
band(sl, 6.32, 0.78, "ALSO CHANGED",
     "Presentation: answers are placed where each question arises. Corrections to v7 are listed in the speaker notes of this slide.", dark=False, size=12.8)
notes(sl, [
 "Reviewer notes transcribed (handwritten sheet): What kind of attack and by whom? References? Does it only work on O-RAN or any other architecture? From which model did we get the KPI? Where did we get the dataset? Who does these attacks and what are the impacts? Again, who does these attacks and how? Were the 8 attacks we considered already existing or did we make them up; if existing, where did we refer them from? Citation of the attack references is really important. Explain the DU layers and the architecture; what kind of attacks and who does it? Should have presented in a better way. Which module will these attacks affect? We need to classify the KPIs so that we can proceed with recovery. Source of attack?",
 "Corrections to v7 carried into v8 (each verified against project files on 3 Oct 2026):",
 "1. Captures: raw PCAPs are not in the packaged archive; 'every frame recorded unmodified' wording replaced on the testbed, pipeline and dataset slides.",
 "2. pmc management telemetry returned no data in any run (not only missing in 48 runs).",
 "3. free_running 1 is set on every daemon: offset measured, clock never steered.",
 "4. A1 wording: a superior data set under the alternate BMCA, not priority2 alone.",
 "5. Timing slide: ±1.5 us is at reference point E against a common time standard.",
 "6. B_bc_replacement cause: the replacement identity is provisioned; the base rule's rogue-BC clause ignores it. Fix updated on the remaining-work slide.",
 "7. Corrective action: none in the 168-run campaign; the 13 Sep pilot executed a port failover (5/5 vs 0/5).",
 "8. Rule v2, frozen before the campaign began, gives verdicts identical to v3 on all 168 runs."])

# ---------------------------------------------------------------- post-build layout fixes (v7-original slides)
def shape_by_text(slide, prefix):
    hits = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith(prefix)]
    assert len(hits) == 1, prefix
    return hits[0]
# v7 S7 "Research gaps": right panel titles/bodies overlapped; re-space five pairs inside the 1.95-6.05 in panel
sl = prs.slides[pos[("v", 7)] - 1]
pairs = ["Matched scenarios", "Two decision arms", "Held-out evaluation", "Trivial control", "Explicit abstention"]
panel = [sh for sh in sl.shapes if not sh.has_text_frame or sh.text_frame.text == ""]
panel = [sh for sh in panel if abs(sh.left - Inches(8.7)) < Inches(0.05) and abs(sh.top - Inches(1.95)) < Inches(0.05)]
assert len(panel) == 1
panel[0].height = Inches(4.25)          # 1.95 -> 6.20 in, still clear of the gap band below
shape_by_text(sl, "RESULTING STUDY DESIGN").top = Inches(2.1)
yy = 2.4
for t in pairs:
    hd = shape_by_text(sl, t)
    idx = list(sl.shapes).index(hd)
    body = list(sl.shapes)[idx + 1]
    bh = 0.56 if t == "Two decision arms" else 0.4
    hd.top = Inches(yy); hd.height = Inches(0.25)
    body.top = Inches(yy + 0.25); body.height = Inches(bh)
    for p in hd.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(14.5)
    for p in body.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(11.5)
    yy += 0.25 + bh + 0.06
# v7 S26 "Measured detection boundary": threshold band body overflowed its band
sl = prs.slides[pos[("v", 26)] - 1]
b = shape_by_text(sl, "At 60 % the ratio")
for p in b.text_frame.paragraphs:
    for r in p.runs: r.font.size = Pt(11.5)
b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
for n, h, o in log: print(f"  v7 S{n}: {h} hit(s)  {o}")
