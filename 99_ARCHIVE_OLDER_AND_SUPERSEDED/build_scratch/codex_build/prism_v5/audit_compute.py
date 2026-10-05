from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import statistics
import zipfile
import importlib.util
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET


ROOT = Path(r"C:\Users\Admin\Documents\AI-Native Self-Healing O-RAN Network using a Digital Twin")
EXTRACTED = Path(r"C:\Users\Admin\AppData\Local\Temp\codex_prism_v5_audit")
REPO = ROOT / "01_CURRENT_SPlane_SelfHealing" / "oran_splane_selfhealing"
ML_PATH = REPO / "ml_comparison_output" / "ML_VS_RULE_COMPARISON.json"
DECK = ROOT / "deliverables" / "ORAN_SPlane_PRISM_Review_v4.pptx"
OUT_DIR = ROOT / ".codex_build" / "prism_v5"

SCENARIOS = [
    "baseline", "A1_rogue_master", "A2_sync_spoof", "A3_replay", "A5_dos_flood",
    "A8_rogue_bc", "B2_gm_failover", "B_bc_replacement", "B3_pdv_congestion",
    "B7_topology_change", "B_unplanned_failover", "C1_removal", "C2_malformed",
    "C3_wholesecond",
]
EXPECTED = {
    "baseline": "BENIGN", "A1_rogue_master": "ATTACK", "A2_sync_spoof": "ATTACK",
    "A3_replay": "ATTACK", "A5_dos_flood": "ATTACK", "A8_rogue_bc": "ATTACK",
    "B2_gm_failover": "BENIGN", "B_bc_replacement": "BENIGN", "B3_pdv_congestion": "BENIGN",
    "B7_topology_change": "BENIGN", "B_unplanned_failover": "UNKNOWN",
    "C1_removal": "ATTACK", "C2_malformed": "ATTACK", "C3_wholesecond": "ATTACK",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def wilson(k: int, n: int, z: float = 1.959963984540054) -> tuple[float, float]:
    p = k / n
    den = 1 + z * z / n
    center = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return center - half, center + half


def decision_metrics():
    rp_spec = importlib.util.spec_from_file_location("randparams", EXTRACTED / "run" / "randparams.py")
    rp = importlib.util.module_from_spec(rp_spec)
    rp_spec.loader.exec_module(rp)
    by_scenario = {}
    all_rows = []
    for scenario in SCENARIOS:
        counts = {"base": Counter(), "v3": Counter()}
        rows = []
        for rep in range(1, 13):
            d = EXTRACTED / "cap" / f"{scenario}__r{rep}"
            ctx = json.loads((d / "context.json").read_text(encoding="utf-8"))
            base = json.loads((d / "decision.json").read_text(encoding="utf-8"))
            v3 = json.loads((d / "decision_v3.json").read_text(encoding="utf-8"))
            counts["base"][base["verdict"]] += 1
            counts["v3"][v3["verdict"]] += 1
            row = {
                "scenario": scenario, "rep": rep, "expected": EXPECTED[scenario],
                "base": base["verdict"], "v3": v3["verdict"],
                "base_fault_hint": base.get("fault_hint"), "v3_fault_hint": v3.get("fault_hint"),
                "truth_fault": ctx.get("fault_id"), "context_rep": ctx.get("rep"),
                "params": rp.params(rep),
            }
            rows.append(row)
            all_rows.append(row)
        by_scenario[scenario] = {"counts": counts, "rows": rows}

    attack_scenarios = [s for s in SCENARIOS if EXPECTED[s] == "ATTACK"]
    benign_scenarios = [s for s in SCENARIOS if EXPECTED[s] == "BENIGN"]
    metrics = {}
    for arm in ("base", "v3"):
        recalls = {s: by_scenario[s]["counts"][arm]["ATTACK"] / 12 for s in attack_scenarios}
        specs = {s: by_scenario[s]["counts"][arm]["BENIGN"] / 12 for s in benign_scenarios}
        metrics[arm] = {
            "recalls": recalls,
            "specificities": specs,
            "macro_sensitivity": statistics.mean(recalls.values()),
            "macro_specificity": statistics.mean(specs.values()),
            "macro_specificity_excluding_bc_replacement": statistics.mean(v for k, v in specs.items() if k != "B_bc_replacement"),
            "attack_correct": sum(by_scenario[s]["counts"][arm]["ATTACK"] for s in attack_scenarios),
            "attack_total": 12 * len(attack_scenarios),
            "benign_correct": sum(by_scenario[s]["counts"][arm]["BENIGN"] for s in benign_scenarios),
            "benign_total": 12 * len(benign_scenarios),
        }

    attack_rows = [r for r in all_rows if r["expected"] == "ATTACK"]
    alias = {"A_intercept": "C1", "A_malformed": "C2", "A_wholesecond": "C3"}
    def attr_correct(r, arm):
        hint = alias.get(r[f"{arm}_fault_hint"], r[f"{arm}_fault_hint"])
        return hint == r["truth_fault"]
    metrics["attribution"] = {
        arm: {"correct": sum(attr_correct(r, arm) for r in attack_rows), "total": len(attack_rows)}
        for arm in ("base", "v3")
    }
    v2v3_equal = True
    for r in all_rows:
        d = EXTRACTED / "cap" / f"{r['scenario']}__r{r['rep']}"
        v2v3_equal &= json.loads((d / "decision_v2.json").read_text(encoding="utf-8"))["verdict"] == r["v3"]
    metrics["regression"] = {
        "v2_v3_identical": v2v3_equal,
        "base_attack_scenario_total": sum(1 for r in all_rows if r["base"] == "ATTACK" and r["expected"] == "ATTACK"),
        "all_decisions_identical_count": sum(r["base"] == r["v3"] for r in all_rows),
    }
    return by_scenario, all_rows, metrics


def c1_metrics(by_scenario):
    records = []
    for row in by_scenario["C1_removal"]["rows"]:
        rep = row["rep"]
        pct = int(row["params"].get("c1_down_pct"))
        p = EXTRACTED / "cap" / f"C1_removal__r{rep}" / "C1_removal.dn.deep.csv"
        data = list(csv.DictReader(p.open(encoding="utf-8-sig", newline="")))
        ts = [int(r["capture_ts_ns"]) for r in data if r.get("capture_ts_ns")]
        span = (max(ts) - min(ts)) / 1e9
        by_msg = defaultdict(list)
        for r in data:
            if r.get("message_type") in {"Sync", "Announce"}:
                by_msg[(r.get("source_clock_identity"), r.get("message_type"))].append(r)
        msg_rows = []
        for (src, msg), group in sorted(by_msg.items()):
            intervals = [int(x["log_message_interval"]) for x in group if x.get("log_message_interval") not in (None, "")]
            med = statistics.median(intervals)
            declared = 2 ** (-med)
            observed = len(group) / span
            msg_rows.append({"src": src, "message": msg, "frames": len(group), "span_s": span,
                             "median_log_interval": med, "declared_hz": declared,
                             "observed_hz": observed, "ratio": observed / declared})
        primary_sync = max((m for m in msg_rows if m["message"] == "Sync"), key=lambda m: m["frames"])
        primary_ann = max((m for m in msg_rows if m["message"] == "Announce"), key=lambda m: m["frames"])
        records.append({"rep": rep, "c1_down_pct": pct, "sync": primary_sync, "announce": primary_ann})
    per_level = {}
    for pct in sorted({r["c1_down_pct"] for r in records}):
        grp = [r for r in records if r["c1_down_pct"] == pct]
        per_level[pct] = {
            "n": len(grp),
            "reps": [r["rep"] for r in grp],
            "sync_median_ratio": statistics.median(r["sync"]["ratio"] for r in grp),
            "sync_mean_ratio": statistics.mean(r["sync"]["ratio"] for r in grp),
            "announce_median_ratio": statistics.median(r["announce"]["ratio"] for r in grp),
            "detected": sum(r["sync"]["ratio"] < 0.5 or r["announce"]["ratio"] < 0.5 for r in grp),
            "run_window": {"duration_s": 44, "warmup_s": 8, "remaining_s": 36,
                           "down_s_floor": math.floor(36 * pct / 100),
                           "tail_s": 36 - math.floor(36 * pct / 100),
                           "available_fraction_of_full_44s": (8 + 36 - math.floor(36 * pct / 100)) / 44},
        }
    levels = sorted(per_level)
    crossing = None
    for a, b in zip(levels, levels[1:]):
        ya = per_level[a]["sync_median_ratio"]
        yb = per_level[b]["sync_median_ratio"]
        if (ya - .5) * (yb - .5) <= 0 and ya != yb:
            crossing = a + (.5 - ya) * (b - a) / (yb - ya)
            break
    return {"records": records, "per_level": per_level, "interpolated_sync_crossing_pct": crossing}


def ml_metrics():
    d = json.loads(ML_PATH.read_text(encoding="utf-8"))
    # Exact McNemar p for two-sided binomial test under p=.5, derived from the 56 paired runs.
    out = {}
    pairs = {
        "Rule_vs_ML": ("armA", "armB"),
        "Rule_vs_OR": ("armA", "combo_or"),
        "Rule_vs_RuleFirst": ("armA", "combo_rule_first_fallback_ml"),
        "Rule_vs_Consensus": ("armA", "combo_consensus_else_abstain"),
        "ML_vs_AlwaysBenign": ("armB", "always_benign"),
    }
    for label, (a, bkey) in pairs.items():
        b = sum(r[a] == r["expected"] and r[bkey] != r["expected"] for r in d["per_run_test_results"])
        c = sum(r[a] != r["expected"] and r[bkey] == r["expected"] for r in d["per_run_test_results"])
        n = b + c
        tail = sum(math.comb(n, i) for i in range(0, min(b, c) + 1)) / 2 ** n if n else 1
        p = min(1.0, 2 * tail)
        out[label] = {"a_correct_b_wrong": b, "a_wrong_b_correct": c, "discordant": n, "exact_two_sided_p": p}
    return d, out


def servo_metrics():
    vals = []
    by_scenario = defaultdict(list)
    pat = re.compile(r"master offset\s+(-?\d+)")
    logs_root = REPO / "ml_comparison_input" / "extracted_168run_ptp4l_logs"
    for p in logs_root.rglob("*.log"):
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            m = pat.search(line)
            if m:
                val = abs(int(m.group(1)))
                vals.append(val)
                scenario = p.parent.name.rsplit("__r", 1)[0]
                by_scenario[scenario].append(val)
    return {
        "parsed_samples": len(vals),
        "mean_abs_ns": statistics.mean(vals) if vals else None,
        "rms_ns": math.sqrt(statistics.mean(v * v for v in vals)) if vals else None,
        "median_abs_ns": statistics.median(vals) if vals else None,
        "selected_scenarios": {
            s: {"parsed_samples": len(by_scenario[s]), "mean_abs_ns": statistics.mean(by_scenario[s]),
                "mean_samples_per_run": len(by_scenario[s]) / 12}
            for s in ("baseline", "C1_removal", "C3_wholesecond")
        },
    }


def deck_text_and_notes():
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    slides = []
    with zipfile.ZipFile(DECK) as z:
        names = set(z.namelist())
        slide_names = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=lambda s: int(re.search(r"\d+", s).group()))
        for i, name in enumerate(slide_names, 1):
            root = ET.fromstring(z.read(name))
            text = [n.text or "" for n in root.findall(".//a:t", ns)]
            notes_name = f"ppt/notesSlides/notesSlide{i}.xml"
            notes = []
            if notes_name in names:
                nroot = ET.fromstring(z.read(notes_name))
                notes = [n.text or "" for n in nroot.findall(".//a:t", ns)]
            slides.append({"slide": i, "text": text, "notes": notes})
    lines = []
    for s in slides:
        lines += [f"# Slide {s['slide']}", "", "## Text", "", "\n".join(s["text"]), "", "## Notes", "", "\n".join(s["notes"]), ""]
    (OUT_DIR / "deck_v4_text_and_notes.md").write_text("\n".join(lines), encoding="utf-8")
    return slides


def workbook_inspect():
    try:
        from openpyxl import load_workbook
    except Exception as e:
        return {"error": repr(e)}
    p = ROOT / "docs" / "comparison_notes_and_results.xlsx"
    if not p.exists():
        return {"path": str(p), "error": "source workbook not found"}
    wb = load_workbook(p, data_only=False, read_only=False)
    data = {"path": str(p), "sheets": {}}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            if any(v is not None for v in vals):
                rows.append(vals)
        data["sheets"][ws.title] = {
            "max_row": ws.max_row, "max_column": ws.max_column,
            "nonempty_rows": rows,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        }
    return data


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    by_scenario, all_rows, metrics = decision_metrics()
    c1 = c1_metrics(by_scenario)
    ml, mcnemar = ml_metrics()
    slides = deck_text_and_notes()
    workbook = workbook_inspect()
    archive = REPO / "gap_coverage_2026-09-20" / "corrected_final" / "splane_campaign_CORRECTED_2026-09-20.tgz"
    required_files = ["context.json", "ptp.pcap", "ptp.deep.csv", "decision.json", "decision_v3.json"]
    run_dirs = [p for p in (EXTRACTED / "cap").iterdir() if p.is_dir()]
    missing = {p.name: [f for f in required_files if not (p / f).exists()] for p in run_dirs}
    missing = {k: v for k, v in missing.items() if v}
    result = {
        "archive": {"path": str(archive), "sha256": sha256(archive), "entries": 1096,
                    "run_dirs": len(run_dirs), "missing_required": missing},
        "per_scenario_counts": {s: {arm: dict(v["counts"][arm]) for arm in ("base", "v3")} for s, v in by_scenario.items()},
        "decision_metrics": metrics,
        "wilson": {"84_of_96": wilson(84, 96), "36_of_60": wilson(36, 60), "72_of_96": wilson(72, 96)},
        "c1": c1,
        "ml": ml,
        "mcnemar_recomputed": mcnemar,
        "servo": servo_metrics(),
        "deck": {"slide_count": len(slides), "slides": slides},
        "workbook": workbook,
    }
    (OUT_DIR / "computed.json").write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    print(json.dumps({
        "archive": result["archive"], "metrics": metrics, "wilson": result["wilson"],
        "c1_per_level": c1["per_level"], "c1_crossing": c1["interpolated_sync_crossing_pct"],
        "ml_keys": {k: ml[k] for k in ["n_runs_found", "required_file_missing_count", "context_rep_mismatch_count", "n_windows", "per_window_metrics", "macro_sensitivity", "macro_specificity", "abstention", "arm_overall", "trivial_baseline_control", "combination_analysis", "fault_class_crossovers", "feature_columns_configured", "feature_columns_populated", "feature_columns_empty", "feature_importances_top15"]},
        "mcnemar": mcnemar, "servo": result["servo"], "deck_slides": len(slides), "workbook_sheets": list(workbook.get("sheets", {})),
    }, indent=2, default=str))


if __name__ == "__main__":
    main()
